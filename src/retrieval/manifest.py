"""Excel của team 3 -> manifest.parquet (một dòng cho mỗi URL ảnh được giữ)."""
import hashlib
import sys
from pathlib import Path

import openpyxl
import pandas as pd

SUMMARY_KEY, APPRAISED_KEY = "Summary", "thẩm định"
KEEP, DROP = "keep", "x"


def _files(data_dir: Path) -> list[Path]:
    """File final trước, recheck sau để recheck ghi đè nhãn. Bỏ file khoá ~$."""
    out: list[Path] = []
    for sub in ("final", "recheck"):
        out += sorted(p for p in (data_dir / sub).glob("*.xlsx") if not p.name.startswith("~$"))
    return out


def _rows(ws):
    """Đọc sheet, đệm dòng bị cắt ô cuối (file không có <dimension>) cho đủ độ dài header."""
    it = ws.iter_rows(values_only=True)
    header = next(it, None)
    if not header:
        return None, iter(())
    n = len(header)
    return list(header), (tuple(r) + (None,) * (n - len(r)) for r in it)


def _pick(sheetnames, key):
    return next((s for s in sheetnames if key in s), None)


def _clean(v):
    return None if v is None or (isinstance(v, str) and not v.strip()) else v


def build_manifest(data_dir: Path) -> pd.DataFrame:
    data_dir = Path(data_dir)
    entities: dict[str, dict] = {}
    pairs: dict[tuple[str, str], dict] = {}  # (entity_id, url) -> trường + nhãn
    for f in _files(data_dir):
        wb = openpyxl.load_workbook(f, read_only=True)
        s = _pick(wb.sheetnames, SUMMARY_KEY)
        if s:
            h, rows = _rows(wb[s])
            if h and "Entity ID" in h:
                ix = {c: h.index(c) for c in h}
                for r in rows:
                    eid = r[ix["Entity ID"]]
                    if eid:
                        entities[eid] = {
                            "entity_name": r[ix["Tên Thực thể"]],
                            "branch": r[ix.get("Nhánh (Branch)", 0)] if "Nhánh (Branch)" in ix else None,
                            "region": r[ix["Địa bàn"]] if "Địa bàn" in ix else None,
                            "culture_type": r[ix["Thể loại Văn hóa"]] if "Thể loại Văn hóa" in ix else None,
                            "entity_desc": r[ix["Mô tả sơ qua về thực thể"]]
                            if "Mô tả sơ qua về thực thể" in ix else None,
                        }
        a = _pick(wb.sheetnames, APPRAISED_KEY)
        if not a:
            continue
        h, rows = _rows(wb[a])
        if not h or "URL Ảnh" not in h:
            continue
        ix = {c: h.index(c) for c in h}
        for r in rows:
            url, eid = _clean(r[ix["URL Ảnh"]]), r[ix["Entity ID"]]
            if not url or not eid:
                continue
            label = DROP if str(r[ix["Giữ/Bỏ"]] or "").strip().lower() == DROP else KEEP
            rec = pairs.setdefault((eid, url), {})
            rec["label"] = label
            rec["entity_name"] = _clean(r[ix.get("Tên Thực thể", 0)]) if "Tên Thực thể" in ix else None
            for col, key in (("Mô tả ảnh (Gemini Vision)", "caption"), ("Tiêu đề trang nguồn", "page_title"),
                             ("Tên miền", "domain"), ("Độ phân giải", "resolution"),
                             ("pHash", "phash"), ("dHash", "dhash")):
                v = _clean(r[ix[col]]) if col in ix else None
                if v is not None or key not in rec:
                    rec[key] = v

    by_url: dict[str, dict] = {}
    for (eid, url), rec in pairs.items():
        if rec["label"] != KEEP:
            continue
        item = by_url.setdefault(url, {"entity_ids": [], "entity_names": [], **{
            k: rec.get(k) for k in ("caption", "page_title", "domain", "resolution", "phash", "dhash")}})
        meta = entities.get(eid, {})
        item["entity_ids"].append(eid)
        item["entity_names"].append(meta.get("entity_name") or rec.get("entity_name"))
        if len(item["entity_ids"]) == 1:  # metadata thực thể lấy theo thực thể đầu tiên
            for k in ("branch", "region", "culture_type", "entity_desc"):
                item[k] = meta.get(k)
    rows_out = [
        {"image_id": hashlib.sha1(u.encode()).hexdigest()[:16], "url": u, **v} for u, v in by_url.items()
    ]
    cols = ["image_id", "url", "entity_ids", "entity_names", "caption", "page_title", "domain",
            "resolution", "phash", "dhash", "branch", "region", "culture_type", "entity_desc"]
    return pd.DataFrame(rows_out, columns=cols)


def write_manifest(df: pd.DataFrame, path: Path) -> None:
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    df.to_parquet(path, index=False)


if __name__ == "__main__":
    src, dst = Path(sys.argv[1]), Path(sys.argv[2])
    manifest = build_manifest(src)
    write_manifest(manifest, dst)
    print(f"rows={len(manifest)}")
